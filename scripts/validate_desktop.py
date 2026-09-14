"""Run a synthetic real-model desktop-controller acceptance cycle.

Use a new output directory and a local model. This does not automate Excel or
certify native Windows behavior. It never reads real transcripts or uploads data.
"""

from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

from transcript_anonymizer import workflow
from transcript_anonymizer.desktop import DesktopController
from transcript_anonymizer.documents import read_docx, validate_output, write_docx


def finish(controller, *, expect_error=False):
    deadline = time.monotonic() + 300
    events = []
    while time.monotonic() < deadline:
        events.extend(controller.messages())
        if any(event["kind"] == "finished" for event in events):
            controller.join(2)
            assert not controller.busy, "Worker did not exit"
            errors = [event for event in events if event["kind"] == "error"]
            assert bool(errors) == expect_error, "Unexpected operation outcome"
            if not expect_error:
                assert any(event["kind"] == "result" for event in events), "No result"
            return events
        time.sleep(0.05)
    controller.cancel()
    controller.join(10)
    raise AssertionError("Synthetic desktop operation timed out")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--require-network-isolation", action="store_true")
    args = parser.parse_args()
    # sysfs can remain host-mounted in an unshare network namespace. The
    # procfs network view is namespace-scoped and therefore the reliable
    # check for this bounded offline validation harness.
    try:
        interfaces = {
            line.split(":", 1)[0].strip()
            for line in Path("/proc/net/dev").read_text(encoding="ascii").splitlines()
            if ":" in line
        }
    except (OSError, UnicodeError):
        interfaces = set()
    isolated = sys.platform == "linux" and interfaces <= {"lo"} and "lo" in interfaces
    if args.require_network_isolation and not isolated:
        parser.error("Run in an isolated network namespace with only loopback")
    args.out.mkdir(parents=True, exist_ok=False)
    started = time.monotonic()
    source = args.out / "Synthetic transcript ä.docx"
    write_docx([
        {"id": "s000001", "text": "00:00:01 Anna Beispiel: Kontakt anna@example.invalid."},
        {"id": "s000002", "text": "00:00:05 Max Muster: Ab 100000 Euro folgt eine Freigabe."},
        {"id": "s000003", "text": "Anna Beispiel: Die Freigabe erfolgt im System."},
    ], source)
    original_hash = workflow.file_hash(source)
    controller = DesktopController(model_path=args.model)
    controller.start_process(source, args.out / "Review folder ü")
    preparation = finish(controller)
    assert controller.status["detector"]["rules_only"] is False
    state, _ = workflow._load(controller.run_path)
    model_findings = sum(row["detector"].startswith("model:") for row in state["findings"])
    assert model_findings > 0, "No real-model findings in the synthetic fixture"
    assert not controller.status["signed_off"]
    controller.export(args.out / "Exports")
    finish(controller, expect_error=True)
    assert controller.export_path is None

    book = load_workbook(controller.status["workbook"])
    sheet = book["Replacements"]
    cell = next(cell for row in sheet for cell in row if cell.value == "Signed off by")
    sheet.cell(cell.row, cell.column + 1).value = "Synthetic reviewer"
    signed = args.out / "Reviewed workbook.xlsx"
    book.save(signed)
    book.close()
    controller.import_workbook(signed)
    finish(controller)
    assert controller.status["signed_off"]
    controller.export(args.out / "Exports")
    finish(controller)
    output = controller.export_path
    assert {path.name for path in output.iterdir()} == {"transcript.docx", "manifest.json"}
    validate_output(output / "transcript.docx")
    text = " ".join(row["text"] for row in read_docx(output / "transcript.docx")["segments"])
    assert all(raw not in text for raw in ("Anna Beispiel", "Max Muster", "anna@example.invalid"))
    assert workflow.file_hash(source) == original_hash
    report = {
        "ok": True,
        "mode": "real-model desktop controller; synthetic input",
        "network_namespace_isolated": isolated,
        "elapsed_s": round(time.monotonic() - started, 2),
        "segments": controller.status["segments"],
        "model_findings": model_findings,
        "progress_events": sum(event["kind"] == "progress" for event in preparation),
        "unsigned_export_blocked": True,
        "signed_export_validated": True,
        "source_unchanged": True,
        "native_windows_acceptance": False,
    }
    print(json.dumps(report, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
