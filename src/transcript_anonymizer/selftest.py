"""Opt-in frozen synthetic workflow smoke. Never consumes user transcripts."""
import json
import socket
from pathlib import Path

from openpyxl import load_workbook

from . import workflow
from .cli import offline_environment
from .detection import Detector, load_policy
from .documents import read_docx, validate_output, write_docx


def run(output: Path, model: Path | None = None, *, detector=None):
    offline_environment()
    output.mkdir(parents=True, exist_ok=False)
    attempts = []
    original = socket.socket.connect
    original_ex = socket.socket.connect_ex
    def blocked(*args, **kwargs):
        attempts.append(True)
        raise OSError('Network is blocked during the synthetic test')
    socket.socket.connect = socket.socket.connect_ex = blocked
    try:
        source = output / 'Synthetic input ü.docx'
        write_docx([
            {'id': 's1', 'text': 'Anna Beispiel: Kontakt anna@example.invalid.'},
            {'id': 's2', 'text': '[00:01:23] Account Manager: Wir verwenden SAP.'},
        ], source)
        before = workflow.file_hash(source)
        policy = load_policy(None)
        detector = detector or Detector(policy, model)
        result = workflow.prepare(source, output / 'review', policy, detector)
        run_path = output / 'review'
        from .errors import AnonymizerError
        try:
            workflow.export(run_path, output / 'unsigned')
        except AnonymizerError:
            pass
        else:
            raise AssertionError('Unsigned export succeeded')
        book = load_workbook(result['workbook'])
        book['Replacements']['B6'] = 'Synthetic reviewer'
        signed = output / 'signed.xlsx'
        book.save(signed)
        book.close()
        result = workflow.import_review(run_path, signed, detector)
        assert result['signed_off']
        book = load_workbook(result['workbook'])
        # Default first group becomes a content-changing removal.
        sheet = book['Replacements']
        columns = {c.value: c.column for c in sheet[9]}
        sheet.cell(10, columns['action']).value = 'remove'
        book.save(output / 'corrected.xlsx')
        book.close()
        result = workflow.import_review(run_path, output / 'corrected.xlsx', detector)
        assert not result['signed_off']
        book = load_workbook(result['workbook'])
        book['Replacements']['B6'] = 'Synthetic reviewer'
        book.save(output / 'resigned.xlsx')
        book.close()
        workflow.import_review(run_path, output / 'resigned.xlsx', detector)
        workflow.export(run_path, output / 'export')
        validate_output(output / 'export/transcript.docx')
        assert {p.name for p in (output / 'export').iterdir()} == {'transcript.docx', 'manifest.json'}
        text = ' '.join(s['text'] for s in read_docx(output / 'export/transcript.docx')['segments'])
        assert 'anna@example.invalid' not in text and 'Anna Beispiel' not in text
        assert '[00:01:23]' in text
        assert workflow.file_hash(source) == before and not attempts
        state, _ = workflow._load(run_path)
        report = {'ok': True, 'model_enabled': not detector.metadata['rules_only'],
                  'model_findings': sum(f['detector'].startswith('model:') for f in state['findings']),
                  'source_unchanged': True, 'network_attempts': 0,
                  'manual_verification': False}
        (output / 'result.json').write_text(json.dumps(report), encoding='utf-8')
        return report
    finally:
        socket.socket.connect = original
        socket.socket.connect_ex = original_ex
