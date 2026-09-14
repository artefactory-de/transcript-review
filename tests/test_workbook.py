from pathlib import Path

import pytest
from openpyxl import load_workbook

from transcript_anonymizer.errors import AnonymizerError
from transcript_anonymizer.workbook import MAX_CELL_CHARS, read_workbook, write_workbook


def view():
    return {
        "schema_version": 1,
        "run_id": "run-opaque-1",
        "revision": 3,
        "candidate_sha256": "a" * 64,
        "signed_off_by": "",
        "groups": [
            {
                "id": "g-1", "source": "Ada Example", "aliases": "Ada",
                "category": "name", "token": "PERSON_1", "count": 2,
                "context": "Ada spoke", "action": "accept", "replacement": "", "note": "",
            },
            {
                "id": "g-2", "source": "+49 123", "aliases": "", "category": "phone",
                "token": "PHONE_1", "count": 1, "context": "call", "action": "accept",
                "replacement": "", "note": "",
            },
        ],
        "passages": [
            {
                "id": "p-1", "source": "Ada spoke", "text": "PERSON_1 spoke", "score": 0.91,
                "reasons": "name signal", "outcome": "uninspected", "action": "keep",
                "replacement": "", "note": "",
            },
            {
                "id": "p-2", "source": "ordinary process", "text": "ordinary process", "score": 0.1,
                "reasons": "", "outcome": "uninspected", "action": "keep", "replacement": "", "note": "",
            },
        ],
        "occurrences": [
            {
                "id": "o-1", "group_id": "g-1", "segment_id": "s-1", "source": "Ada",
                "context": "Ada spoke", "action": "inherit", "replacement": "", "note": "",
            }
        ],
    }


def test_round_trip_protection_and_editable_cells(tmp_path: Path):
    expected = view()
    path = tmp_path / "review.xlsx"
    write_workbook(expected, path)

    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["Start here", "Replacements", "Ranked passages", "Occurrences"]
    assert workbook.security.lockStructure
    assert workbook["Replacements"].protection.sheet
    assert workbook["Replacements"]["A10"].protection.locked
    assert not workbook["Replacements"]["H10"].protection.locked
    assert not workbook["Replacements"]["B6"].protection.locked
    assert workbook["Ranked passages"]["B10"].protection.locked
    assert not workbook["Ranked passages"]["F10"].protection.locked
    assert not workbook["Occurrences"]["F10"].protection.locked
    assert workbook["Replacements"].data_validations.count == 1
    assert workbook["Ranked passages"].data_validations.count == 2
    from openpyxl.cell.cell import Cell

    for sheet in workbook:
        assert sheet.protection.sheet
        assert sheet.protection.formatColumns is False
        assert sheet.protection.formatCells is True
        assert sheet.protection.insertColumns is True
        assert sheet.protection.deleteColumns is True
        for row in sheet:
            for cell in row:
                if isinstance(cell, Cell):
                    assert cell.alignment.vertical == "top", (sheet.title, cell.coordinate)
    workbook.close()

    imported = read_workbook(path, expected)
    assert imported["groups"] == expected["groups"]
    assert imported["passages"] == expected["passages"]
    assert imported["signed_off_by"] == ""


def test_column_width_changes_keep_protection_and_allow_review_import(tmp_path):
    expected = view()
    path = tmp_path / 'resized.xlsx'
    write_workbook(expected, path)
    book = load_workbook(path)
    for sheet in book:
        sheet.column_dimensions['B'].width = 70
    book.save(path)
    book.close()
    reopened = load_workbook(path)
    assert reopened['Ranked passages'].column_dimensions['B'].width == 70
    assert reopened['Ranked passages']['B10'].protection.locked
    assert not reopened['Ranked passages']['H10'].protection.locked
    assert reopened['Ranked passages'].protection.sheet
    reopened.close()
    assert read_workbook(path, expected)['passages'] == expected['passages']


def test_in_workbook_guidance_covers_choices_and_stays_protected(tmp_path):
    from transcript_anonymizer import workbook as module

    expected = view() | {"source_filename": "Workshop März.docx"}
    path = tmp_path / "instructions.xlsx"
    write_workbook(expected, path)
    book = load_workbook(path)
    assert book.active.title == "Start here"
    guide = book["Start here"]
    assert guide["B2"].value == "Workshop März.docx"
    assert guide.protection.sheet
    assert all(cell.protection.locked for row in guide for cell in row)
    text = "\n".join(str(cell.value) for row in guide for cell in row if cell.value)
    assert tuple(module._DECISION_HELP[("Replacements", "action")]) == module._GROUP_ACTIONS
    assert tuple(module._DECISION_HELP[("Ranked passages", "outcome")]) == module._PASSAGE_OUTCOMES
    assert tuple(module._DECISION_HELP[("Ranked passages", "action")]) == module._PASSAGE_ACTIONS
    assert tuple(module._DECISION_HELP[("Occurrences", "action")]) == module._OCCURRENCE_ACTIONS
    for (sheet, field), meanings in module._DECISION_HELP.items():
        assert f"{sheet}: {field}" in text
        for value, meaning in meanings.items():
            assert value in text and meaning in text
        ws = book[sheet]
        header = next(cell for cell in ws[9] if cell.value == field)
        assert all(value in header.comment.text for value in meanings)
    assert "ENTIRE passage" in text
    assert "Content changes clear sign-off" in text
    assert "Ctrl+C" in text and "Paste > Values" in text
    assert "Do not copy source" in text
    assert "several items in this one cell" in text
    assert "Example: two missed items" in text
    assert "existing [PERSON_1] token stay unchanged" in text
    assert "They are identical when no change" in text
    assert "unread" in text and "false positive" in text
    for sheet in ("Replacements", "Ranked passages", "Occurrences"):
        ws = book[sheet]
        assert "Start here" in ws["A7"].value
        assert ws["A8"].value and ws["A8"].protection.locked
        for validation in ws.data_validations.dataValidation:
            assert 0 < len(validation.prompt) <= 255
            assert "Select a value from the list." != validation.prompt
    book.save(path)
    book.close()
    assert read_workbook(path, expected)["groups"] == expected["groups"]


def test_legacy_workbook_without_guidance_still_imports(tmp_path):
    path = tmp_path / "legacy.xlsx"
    write_workbook(view(), path)
    book = load_workbook(path)
    del book["Start here"]
    book.save(path)
    book.close()
    assert read_workbook(path, view())["groups"] == view()["groups"]


@pytest.mark.parametrize("invalid", ["formula", "extra_sheet"])
def test_guidance_sheet_cannot_smuggle_formulas_or_unknown_sheets(tmp_path, invalid):
    path = tmp_path / "unsafe-guide.xlsx"
    write_workbook(view(), path)
    book = load_workbook(path)
    if invalid == "formula":
        book["Start here"]["B2"] = "=1+1"
    else:
        book.create_sheet("Extra instructions")
    book.save(path)
    book.close()
    with pytest.raises(AnonymizerError, match="formula|sheets"):
        read_workbook(path, view())


def test_guidance_without_findings_is_readable_and_importable(tmp_path):
    expected = view() | {"groups": [], "passages": [], "occurrences": []}
    path = tmp_path / "empty.xlsx"
    write_workbook(expected, path)
    book = load_workbook(path)
    assert book.active.title == "Start here"
    assert "Occurrences" not in book.sheetnames
    book.close()
    assert read_workbook(path, expected)["groups"] == []


def test_input_filename_label_is_bound_to_the_run(tmp_path):
    expected = view() | {"source_filename": "Workshop.docx"}
    path = tmp_path / "review.xlsx"
    write_workbook(expected, path)
    book = load_workbook(path)
    book["Start here"]["B2"] = "Other input.docx"
    book.save(path)
    book.close()
    with pytest.raises(AnonymizerError, match="immutable input filename"):
        read_workbook(path, expected)


def test_editable_decisions_and_single_signoff(tmp_path: Path):
    expected = view()
    path = tmp_path / "review.xlsx"
    write_workbook(expected, path)
    workbook = load_workbook(path)
    replacements = workbook["Replacements"]
    replacements["H10"] = "replace"
    replacements["I10"] = "PERSON_REVIEWED"
    replacements["J10"] = "reviewed"
    replacements["B6"] = "Reviewer One"
    replacements["B6"].data_type = "s"
    passages = workbook["Ranked passages"]
    passages["F10"] = "clean"
    passages["G10"] = "keep"
    workbook.save(path)
    workbook.close()

    result = read_workbook(path, expected)
    assert result["signed_off_by"] == "Reviewer One"
    assert result["groups"][0]["action"] == "replace"
    assert result["groups"][0]["replacement"] == "PERSON_REVIEWED"
    assert result["passages"][0]["outcome"] == "clean"


def test_row_sorting_is_permitted_and_result_is_canonical(tmp_path: Path):
    expected = view()
    path = tmp_path / "review.xlsx"
    write_workbook(expected, path)
    workbook = load_workbook(path)
    sheet = workbook["Replacements"]
    values = [[sheet.cell(row, col).value for col in range(1, 11)] for row in (10, 11)]
    for col, value in enumerate(values[0], 1):
        sheet.cell(10, col).value = values[1][col - 1]
        sheet.cell(10, col).data_type = "s" if isinstance(values[1][col - 1], str) else "n"
        sheet.cell(11, col).value = values[0][col - 1]
        sheet.cell(11, col).data_type = "s" if isinstance(values[0][col - 1], str) else "n"
    workbook.save(path)
    workbook.close()
    assert [row["id"] for row in read_workbook(path, expected)["groups"]] == ["g-1", "g-2"]


@pytest.mark.parametrize("mutate", ["immutable", "duplicate", "foreign", "missing", "bad_enum", "formula"])
def test_invalid_imports_are_rejected(tmp_path: Path, mutate: str):
    expected = view()
    path = tmp_path / "review.xlsx"
    write_workbook(expected, path)
    workbook = load_workbook(path)
    sheet = workbook["Replacements"]
    if mutate == "immutable":
        sheet["B10"] = "changed source"
    elif mutate == "duplicate":
        sheet["A11"] = "g-1"
    elif mutate == "foreign":
        sheet["A10"] = "g-foreign"
    elif mutate == "missing":
        for cell in sheet[11]:
            cell.value = None
    elif mutate == "bad_enum":
        sheet["H10"] = "not-an-action"
    else:
        sheet["H10"] = "=1+1"
    workbook.save(path)
    workbook.close()
    with pytest.raises(AnonymizerError) as exc:
        read_workbook(path, expected)
    assert "Ada Example" not in str(exc.value)
    assert "g-foreign" not in str(exc.value)


def test_bypass_protection_and_stale_revision_are_rejected(tmp_path: Path):
    expected = view()
    path = tmp_path / "review.xlsx"
    write_workbook(expected, path)
    workbook = load_workbook(path)
    workbook.security.lockStructure = False
    workbook["Replacements"].protection.sheet = False
    workbook["Replacements"]["C10"] = "tampered alias"
    workbook.save(path)
    workbook.close()
    with pytest.raises(AnonymizerError, match="immutable"):
        read_workbook(path, expected)

    workbook = load_workbook(path)
    workbook["Replacements"]["C10"] = expected["groups"][0]["aliases"]
    workbook["Replacements"]["B4"] = expected["revision"] - 1
    workbook.save(path)
    workbook.close()
    with pytest.raises(AnonymizerError, match="stale"):
        read_workbook(path, expected)


def test_miss_requires_correction(tmp_path: Path):
    expected = view()
    path = tmp_path / "review.xlsx"
    write_workbook(expected, path)
    workbook = load_workbook(path)
    sheet = workbook["Ranked passages"]
    sheet["F10"] = "miss"
    workbook.save(path)
    workbook.close()
    with pytest.raises(AnonymizerError, match="miss requires correction"):
        read_workbook(path, expected)


def test_formula_like_source_is_literal_but_formula_cell_is_rejected(tmp_path: Path):
    expected = view()
    expected["groups"][1]["source"] = "=not a formula"
    path = tmp_path / "review.xlsx"
    write_workbook(expected, path)
    workbook = load_workbook(path, data_only=False)
    assert workbook["Replacements"]["B11"].data_type == "s"
    workbook["Replacements"]["H10"] = "=NOW()"
    workbook.save(path)
    workbook.close()
    with pytest.raises(AnonymizerError, match="formula"):
        read_workbook(path, expected)


def test_oversized_cell_is_blocked_without_truncation(tmp_path: Path):
    expected = view()
    expected["groups"][0]["context"] = "x" * (MAX_CELL_CHARS + 1)
    with pytest.raises(AnonymizerError, match="Excel cell limit"):
        write_workbook(expected, tmp_path / "review.xlsx")


def test_distant_sparse_cell_is_rejected_before_openpyxl_iteration(tmp_path: Path):
    expected = view()
    path = tmp_path / "review.xlsx"
    write_workbook(expected, path)
    workbook = load_workbook(path)
    workbook["Replacements"]["XFD1048576"] = "unexpected"
    workbook.save(path)
    workbook.close()
    with pytest.raises(AnonymizerError, match="worksheet dimensions"):
        read_workbook(path, expected)
