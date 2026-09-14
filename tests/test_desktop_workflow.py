"""Headless release checks for the public desktop controller."""

from __future__ import annotations

import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from transcript_anonymizer import workflow
from transcript_anonymizer.desktop import DesktopController, eta_seconds, format_eta
from transcript_anonymizer.detection import Detector
from transcript_anonymizer.documents import read_docx, write_docx
from transcript_anonymizer.errors import AnonymizerError


def _source(tmp_path: Path, *texts: str) -> Path:
    source = tmp_path / f"source-{len(list(tmp_path.glob('source-*.docx')))}.docx"
    write_docx([{"id": f"s{i:06d}", "text": text} for i, text in enumerate(texts, 1)], source)
    return source


def _rules_factory(policy, _model_path, progress):
    detector = Detector(policy, rules_only=True)
    detector.progress = lambda completed, total: progress("processing", completed, total)
    return detector


def _controller() -> DesktopController:
    return DesktopController(model_path=None, detector_factory=_rules_factory)


def _finish(controller: DesktopController, timeout: float = 5.0) -> list[dict]:
    deadline = time.monotonic() + timeout
    messages: list[dict] = []
    while time.monotonic() < deadline:
        messages.extend(controller.messages())
        if any(message["kind"] == "finished" for message in messages):
            controller.join(0.1)
            messages.extend(controller.messages())
            return messages
        time.sleep(0.01)
    controller.join(0.1)
    messages.extend(controller.messages())
    raise AssertionError(f"controller did not finish: {messages}")


def _sign(workbook: Path, tmp_path: Path, name: str = "Synthetic reviewer") -> Path:
    book = load_workbook(workbook)
    ws = book["Replacements"]
    label = next(cell for row in ws for cell in row if cell.value == "Signed off by")
    ws.cell(label.row, label.column + 1).value = name
    signed = tmp_path / f"signed-{len(list(tmp_path.glob('signed-*.xlsx')))}.xlsx"
    book.save(signed)
    book.close()
    return signed


def _edit_group(workbook: Path, tmp_path: Path, group_id: str, **edits: str) -> Path:
    book = load_workbook(workbook)
    ws = book["Replacements"]
    header = next(row for row in ws if row[0].value == "id")
    columns = {cell.value: cell.column for cell in header}
    row = next(row for row in ws if row[0].value == group_id)
    for field, value in edits.items():
        ws.cell(row[0].row, columns[field]).value = value
    edited = tmp_path / f"edited-{len(list(tmp_path.glob('edited-*.xlsx')))}.xlsx"
    book.save(edited)
    book.close()
    return edited


def test_rules_only_controller_requires_signoff_then_exports(tmp_path: Path):
    controller = _controller()
    source = _source(tmp_path, "Anna Beispiel: anna@example.invalid")
    review_root = tmp_path / "reviews"

    controller.start_process(source, review_root)
    messages = _finish(controller)
    result = next(message["result"] for message in messages if message["kind"] == "result")
    assert result["signed_off"] is False
    assert controller.status["signed_off"] is False

    controller.export(tmp_path / "exports")
    blocked = _finish(controller)
    assert any(message["kind"] == "error" for message in blocked)
    assert not list((tmp_path / "exports").glob("export-*/transcript.docx"))

    controller.import_workbook(_sign(Path(result["workbook"]), tmp_path))
    imported = _finish(controller)
    signed = next(message["result"] for message in imported if message["kind"] == "result")
    assert signed["signed_off"] is True

    controller.export(tmp_path / "exports")
    exported = _finish(controller)
    assert any(message["kind"] == "result" for message in exported)
    output = next(tmp_path.glob("exports/export-*/transcript.docx"))
    text = " ".join(segment["text"] for segment in read_docx(output)["segments"])
    assert "Anna Beispiel" not in text
    assert "anna@example.invalid" not in text


def test_corrected_import_clears_signoff_and_blocks_export(tmp_path: Path):
    controller = _controller()
    source = _source(tmp_path, "Anna Beispiel: Hallo")
    controller.start_process(source, tmp_path / "reviews")
    result = next(message["result"] for message in _finish(controller) if message["kind"] == "result")
    controller.import_workbook(_sign(Path(result["workbook"]), tmp_path))
    signed = next(message["result"] for message in _finish(controller) if message["kind"] == "result")

    state, _ = workflow._load(controller.run_path)
    group_id = state["view"]["groups"][0]["id"]
    corrected_book = _edit_group(
        Path(signed["workbook"]), tmp_path, group_id, action="replace", replacement="PERSON_CUSTOM"
    )
    controller.import_workbook(corrected_book)
    corrected = next(message["result"] for message in _finish(controller) if message["kind"] == "result")
    assert corrected["signed_off"] is False

    controller.export(tmp_path / "exports")
    blocked = _finish(controller)
    assert any(message["kind"] == "error" for message in blocked)


def test_old_persisted_overlap_resolver_state_blocks_controller_export(tmp_path: Path):
    controller = _controller()
    source = _source(tmp_path, "Anna Beispiel")
    controller.start_process(source, tmp_path / "reviews")
    result = next(message["result"] for message in _finish(controller) if message["kind"] == "result")
    controller.import_workbook(_sign(Path(result["workbook"]), tmp_path))
    next(message["result"] for message in _finish(controller) if message["kind"] == "result")
    run = controller.run_path
    state, folder = workflow._load(run)
    state["overlap_resolver_version"] = 1
    raw = workflow.encoded(state)
    (folder / "state.json").write_bytes(raw)
    (run / "current.json").write_bytes(
        workflow.encoded({"directory": folder.name, "sha256": workflow.digest(raw)})
    )

    controller.select_run(run)
    selected = next(message["result"] for message in _finish(controller) if message["kind"] == "result")
    assert selected["signed_off"] is True
    controller.export(tmp_path / "exports")
    blocked = _finish(controller)
    assert any("predates" in message.get("message", "") for message in blocked)
    assert not list((tmp_path / "exports").glob("export-*/transcript.docx"))


def test_busy_rejects_concurrent_model_job_and_cancel_has_no_success(tmp_path: Path):
    started = False

    class SlowDetector:
        def __init__(self, policy, _model_path, progress):
            self.policy = policy
            self.metadata = {"detector": "slow-synthetic", "version": 1}
            self.progress = progress

        def detect(self, segments):
            nonlocal started
            started = True
            for completed in range(len(segments) + 1):
                self.progress("processing", completed, len(segments))
                time.sleep(0.1)
            return []

    controller = DesktopController(detector_factory=SlowDetector)
    source = _source(tmp_path, *[f"segment {i}" for i in range(20)])
    controller.start_process(source, tmp_path / "reviews")
    deadline = time.monotonic() + 2
    while not started and time.monotonic() < deadline:
        time.sleep(0.01)
    with pytest.raises(AnonymizerError, match="still running"):
        controller.start_process(source, tmp_path / "reviews")
    assert controller.cancel() is True
    messages = _finish(controller)
    assert any(message["kind"] == "cancelled" for message in messages)
    assert not any(message["kind"] == "result" for message in messages)
    assert controller.status is None


def test_failed_new_prepare_clears_last_successful_run_state(tmp_path: Path):
    controller = _controller()
    source = _source(tmp_path, "Anna Beispiel")
    controller.start_process(source, tmp_path / "reviews")
    _finish(controller)
    assert controller.status is not None

    controller.start_process(tmp_path / "missing.docx", tmp_path / "reviews")
    failure = _finish(controller)
    assert any(message["kind"] == "error" for message in failure)
    assert controller.status is None
    assert controller.run_path is None

    controller.start_process(source, tmp_path / "reviews")
    recovered = _finish(controller)
    assert any(message["kind"] == "result" for message in recovered)


def test_eta_is_unknown_until_coherent_progress_samples_exist():
    assert eta_seconds([], 1, 3) is None
    assert eta_seconds([(10.0, 1)], 1, 3) is None
    assert eta_seconds([(10.0, 1), (11.0, 2)], 2, 4) is None
    assert eta_seconds([(10.0, 2), (11.0, 1)], 1, 4) is None
    assert eta_seconds([(10.0, 1), (11.0, 2)], 4, 4) is None
    assert format_eta(None) == "Estimating..."
    assert format_eta(62) == "About 2 min remaining"
