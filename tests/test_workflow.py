from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from transcript_anonymizer import workflow
from transcript_anonymizer.detection import Detector, load_policy
from transcript_anonymizer.documents import read_docx, write_docx
from transcript_anonymizer.errors import AnonymizerError


@pytest.fixture
def prepared(tmp_path):
    policy = load_policy(None)
    policy["aliases"] = [{"entity_key": "a", "aliases": ["Anna Beispiel", "A. Beispiel"]}]
    detector = Detector(policy, rules_only=True)
    source = tmp_path / "source.docx"
    write_docx(
        [
            {"id": "s000001", "text": "00:00 Anna Beispiel: anna@example.invalid", "locator": "p1"},
            {
                "id": "s000002",
                "text": "00:01 A. Beispiel: Die Grenze ist 100000 Euro.",
                "locator": "p2",
            },
            {"id": "s000003", "text": "00:02 Die Freigabe wird dokumentiert.", "locator": "p3"},
        ],
        source,
    )
    before = workflow.file_hash(source)
    run = tmp_path / "run"
    result = workflow.prepare(source, run, policy, detector)
    assert workflow.file_hash(source) == before
    return run, detector, result


def edited(result, tmp_path, sheet=None, record_id=None, edits=None, signoff=None):
    book = load_workbook(result["workbook"])
    if signoff is not None:
        main = book["Replacements"]
        target = next(cell for row in main for cell in row if cell.value == "Signed off by")
        main.cell(target.row, target.column + 1, signoff)
    if sheet:
        ws = book[sheet]
        header = next(row for row in ws if row[0].value == "id")
        columns = {cell.value: cell.column for cell in header}
        row = next(row for row in ws if row[0].value == record_id)
        for key, value in edits.items():
            ws.cell(row[0].row, columns[key], value)
    path = tmp_path / f"edit-{len(list(tmp_path.glob('edit-*')))}.xlsx"
    book.save(path)
    return path


def test_signoff_export_unread_and_no_leak(prepared, tmp_path):
    run, _, result = prepared
    assert Path(result["workbook"]).name == "source - review.xlsx"
    with pytest.raises(AnonymizerError, match="sign-off"):
        workflow.export(run, tmp_path / "unsigned")
    assert not (tmp_path / "unsigned").exists()
    workbook = edited(result, tmp_path, signoff="Synthetic reviewer")
    signed = workflow.import_review(run, workbook)
    assert signed["signed_off"]
    assert Path(signed["workbook"]).name == "source - review.xlsx"
    assert signed["uninspected_passages"] == signed["segments"]
    workflow.export(run, tmp_path / "export")
    assert {p.name for p in (tmp_path / "export").iterdir()} == {"transcript.docx", "manifest.json"}
    texts = " ".join(s["text"] for s in read_docx(tmp_path / "export/transcript.docx")["segments"])
    assert "Anna Beispiel" not in texts
    assert "A. Beispiel" not in texts
    assert "anna@example.invalid" not in texts
    assert "100000 Euro" in texts
    assert "Synthetic reviewer" not in (tmp_path / "export/manifest.json").read_text()
    assert "source.docx" not in (tmp_path / "export/manifest.json").read_text()
    assert workflow.file_hash(Path(signed["candidate"])) == workflow.file_hash(
        tmp_path / "export/transcript.docx"
    )
    assert workflow.import_review(run, workbook)["import_result"] == "already_applied"


def test_content_edit_clears_same_import_signoff(prepared, tmp_path):
    run, detector, result = prepared
    state, _ = workflow._load(run)
    gid = state["view"]["groups"][0]["id"]
    workbook = edited(
        result,
        tmp_path,
        "Replacements",
        gid,
        {"action": "replace", "replacement": "PARTICIPANT_X"},
        "Reviewer",
    )
    updated = workflow.import_review(run, workbook, detector)
    assert updated["import_result"] == "candidate_changed_signoff_cleared"
    assert not updated["signed_off"]
    with pytest.raises(AnonymizerError):
        workflow.export(run, tmp_path / "not-approved")
    signed = workflow.import_review(run, edited(updated, tmp_path, signoff="Reviewer"))
    assert signed["signed_off"]


def test_passage_correction_rescanned_and_unchanged_review_preserved(prepared, tmp_path):
    run, detector, result = prepared
    state, _ = workflow._load(run)
    pid = state["view"]["passages"][0]["id"]
    workbook = edited(
        result,
        tmp_path,
        "Ranked passages",
        pid,
        {"action": "replace", "replacement": "Kontakt neu@example.invalid", "outcome": "miss"},
        "Reviewer",
    )
    updated = workflow.import_review(run, workbook, detector)
    candidate = " ".join(s["text"] for s in read_docx(Path(updated["candidate"]))["segments"])
    assert "neu@example.invalid" not in candidate
    assert not updated["signed_off"]


def test_tampered_candidate_stale_workbook_and_atomic_failure(prepared, tmp_path):
    run, _, result = prepared
    original_pointer = (run / "current.json").read_bytes()
    workbook = edited(result, tmp_path, signoff="Reviewer")
    with (
        patch.object(workflow, "write_workbook", side_effect=AnonymizerError("test failure")),
        pytest.raises(AnonymizerError),
    ):
        workflow.import_review(run, workbook)
    assert (run / "current.json").read_bytes() == original_pointer
    assert not (run / ".lock").exists()
    signed = workflow.import_review(run, workbook)
    stale = edited(result, tmp_path, signoff="Different reviewer")
    with pytest.raises(AnonymizerError):
        workflow.import_review(run, stale)
    with Path(signed["candidate"]).open("ab") as handle:
        handle.write(b"tampered")
    with pytest.raises(AnonymizerError, match="Candidate bytes"):
        workflow.export(run, tmp_path / "export")


def test_lock_and_source_destination_safety(prepared, tmp_path):
    run, detector, result = prepared
    with workflow.run_lock(run), pytest.raises(AnonymizerError, match="locked"):
        workflow.import_review(run, Path(result["workbook"]))
    with pytest.raises(AnonymizerError):
        workflow.export(run, run / "export")
    with pytest.raises(AnonymizerError, match="already exists"):
        workflow.prepare(tmp_path / "source.docx", run, load_policy(None), detector)


def test_generated_tokens_do_not_raise_retained_text_risk():
    from transcript_anonymizer.detection import rank_passages

    candidate = {"id": "s1", "text": "PERSON_001 hat zugestimmt."}
    ranked = rank_passages([candidate], [])
    assert "identifier_like_token" not in ranked[0]["reasons"]
    original = {"id": "s2", "text": "Referenz ABC12345."}
    assert "identifier_like_token" in rank_passages([original], [])[0]["reasons"]


@pytest.mark.parametrize("name", ["Workshop März.docx", "Other workshop.docx"])
def test_source_named_workbook_and_identity(tmp_path, name):
    source = tmp_path / name
    write_docx([{"id": "s1", "text": "Routine process."}], source)
    policy = load_policy(None)
    result = workflow.prepare(source, tmp_path / "run", policy, Detector(policy, rules_only=True))
    assert Path(result["workbook"]).name == f"{source.stem} - review.xlsx"
    book = load_workbook(result["workbook"])
    assert book["Start here"]["B2"].value == name
    book.close()


@pytest.mark.parametrize("name", ["CON.docx", "bad:name?.docx", "../file.docx", "ä" * 200 + ".docx"])
def test_workbook_filename_is_bounded_and_safe(name):
    result = workflow._review_filename(Path(name))
    assert not workflow._UNSAFE_FILENAME.search(result)
    assert len(result.encode("utf-8")) < 200
    assert workflow._workbook_filename({"workbook_filename": result}) == result


def test_legacy_filename_and_unsafe_state_names():
    assert workflow._workbook_filename({}) == "review.xlsx"
    for name in ("../review.xlsx", "C:\\review.xlsx", "/review.xlsx", "CON.xlsx", 123):
        with pytest.raises(AnonymizerError, match="filename"):
            workflow._workbook_filename({"workbook_filename": name})


def test_existing_generic_named_run_imports_without_renaming(prepared, tmp_path):
    run, _, result = prepared
    state, folder = workflow._load(run)
    Path(result["workbook"]).rename(folder / "review.xlsx")
    del state["workbook_filename"]
    del state["source_filename"]
    state["view"].pop("source_filename", None)
    raw = workflow.encoded(state)
    (folder / "state.json").write_bytes(raw)
    (run / "current.json").write_bytes(workflow.encoded({
        "directory": folder.name, "sha256": workflow.digest(raw),
    }))
    legacy = workflow.status(run)
    assert Path(legacy["workbook"]).name == "review.xlsx"
    signed = workflow.import_review(run, edited(legacy, tmp_path, signoff="Reviewer"))
    assert signed["signed_off"]
    assert Path(signed["workbook"]).name == "review.xlsx"
    workflow.export(run, tmp_path / "legacy-export")
