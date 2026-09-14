"""Independent overlap and review-state regressions for workflow integration."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook

from transcript_anonymizer import workflow
from transcript_anonymizer.documents import read_docx, write_docx
from transcript_anonymizer.errors import AnonymizerError


class FakeDetector:
    def __init__(self, finder):
        self.policy = {"schema_version": 1}
        self.metadata = {"detector": "synthetic-overlap", "version": 1}
        self.finder = finder
        self.calls: list[list[dict]] = []

    def detect(self, segments):
        snapshot = deepcopy(segments)
        self.calls.append(snapshot)
        return self.finder(snapshot)


def _finding(
    segment,
    start,
    end,
    *,
    score=0.8,
    review_only=False,
    category="person",
    entity_key=None,
):
    row = {
        "id": f"synthetic:{segment['id']}:{start}:{end}:{category}:{score}",
        "segment_id": segment["id"],
        "start": start,
        "end": end,
        "text": segment["text"][start:end],
        "category": category,
        "score": score,
        "detector": "synthetic:overlap",
    }
    if review_only:
        row["review_only"] = True
    if entity_key is not None:
        row["entity_key"] = entity_key
    return row


def _source(tmp_path: Path, texts: list[str]) -> Path:
    source = tmp_path / "source.docx"
    write_docx(
        [{"id": f"s{i:06d}", "text": text} for i, text in enumerate(texts, 1)],
        source,
    )
    return source


def _prepare(tmp_path: Path, texts: list[str], finder):
    detector = FakeDetector(finder)
    source = _source(tmp_path, texts)
    run = tmp_path / "run"
    policy = detector.policy
    result = workflow.prepare(source, run, policy, detector)
    return run, detector, result


def _edit_workbook(
    workbook_path: str | Path,
    tmp_path: Path,
    *,
    sheet: str | None = None,
    record_id: str | None = None,
    edits: dict | None = None,
    signoff: str | None = None,
) -> Path:
    book = load_workbook(workbook_path)
    if signoff is not None:
        ws = book["Replacements"]
        label = next(cell for row in ws for cell in row if cell.value == "Signed off by")
        ws.cell(label.row, label.column + 1).value = signoff
    if sheet is not None:
        ws = book[sheet]
        header_row = next(row for row in ws if row[0].value == "id")
        columns = {cell.value: cell.column for cell in header_row}
        row = next(row for row in ws if row[0].value == record_id)
        for field, value in (edits or {}).items():
            ws.cell(row[0].row, columns[field]).value = value
    output = tmp_path / f"workbook-{len(list(tmp_path.glob('workbook-*.xlsx')))}.xlsx"
    book.save(output)
    book.close()
    return output


def test_full_name_overlap_cannot_be_reduced_to_short_name(tmp_path: Path):
    def finder(segments):
        segment = segments[0]
        return [
            _finding(segment, 0, 13, score=0.78),
            _finding(segment, 0, 4, score=0.99),
        ]

    run, _, _ = _prepare(tmp_path, ["Anna Beispiel arbeitet heute."], finder)
    state, _ = workflow._load(run)

    selected = state["selected"]
    assert any(item["start"] == 0 and item["end"] == 13 for item in selected)
    rendered = workflow._render(state)[0]["text"]
    assert "Beispiel" not in rendered


def test_crossing_and_chained_overlaps_preserve_union_of_source_coverage(tmp_path: Path):
    def finder(segments):
        segment = segments[0]
        return [
            _finding(segment, 0, 4, score=0.8),
            _finding(segment, 3, 7, score=0.8),
            _finding(segment, 6, 10, score=0.8),
        ]

    run, _, _ = _prepare(tmp_path, ["abcdefghij"], finder)
    state, _ = workflow._load(run)
    covered = {
        offset
        for finding in state["selected"]
        for offset in range(finding["start"], finding["end"])
    }
    assert covered == set(range(10))


def test_review_only_evidence_does_not_extend_mask(tmp_path: Path):
    def finder(segments):
        segment = segments[0]
        return [
            _finding(segment, 0, 4, score=0.99),
            _finding(segment, 0, 13, score=0.4, review_only=True),
        ]

    run, _, _ = _prepare(tmp_path, ["Anna Beispiel"], finder)
    state, _ = workflow._load(run)

    assert [(item["start"], item["end"]) for item in state["selected"]] == [(0, 4)]
    assert any(item.get("review_only") for item in state["suppressed"])


def test_best_effort_aliases_share_group_and_export_has_no_source_name(tmp_path: Path):
    def finder(segments):
        return [
            _finding(segments[0], 0, 13, entity_key="person:anna"),
            _finding(segments[1], 0, 13, entity_key="person:anna"),
            _finding(segments[2], 0, 11, entity_key="person:anna"),
        ]

    run, _, result = _prepare(
        tmp_path,
        ["Anna Beispiel: Hallo", "Anna Beispiel meldet sich", "A. Beispiel wartet"],
        finder,
    )
    state, _ = workflow._load(run)
    assert len(state["groups"]) == 1

    workflow.import_review(run, _edit_workbook(result["workbook"], tmp_path, signoff="Reviewer"))
    destination = tmp_path / "export"
    workflow.export(run, destination)
    text = " ".join(segment["text"] for segment in read_docx(destination / "transcript.docx")["segments"])
    assert "Anna Beispiel" not in text


def test_xlsx_signoff_allows_export_only_after_import(tmp_path: Path):
    def finder(segments):
        return [_finding(segments[0], 0, 13)]

    run, _, result = _prepare(tmp_path, ["Anna Beispiel"], finder)
    with pytest.raises(AnonymizerError, match="sign-off"):
        workflow.export(run, tmp_path / "before-signoff")

    signed = workflow.import_review(
        run, _edit_workbook(result["workbook"], tmp_path, signoff="Synthetic reviewer")
    )
    assert signed["signed_off"]
    workflow.export(run, tmp_path / "approved")
    assert (tmp_path / "approved" / "transcript.docx").is_file()


def test_passage_correction_rescans_full_working_document(tmp_path: Path):
    def finder(segments):
        return [_finding(segments[0], 0, 13)] if "Anna Beispiel" in segments[0]["text"] else []

    run, detector, result = _prepare(
        tmp_path, ["Anna Beispiel", "Ordinary review text"], finder
    )
    state, _ = workflow._load(run)
    passage_id = next(row["id"] for row in state["view"]["passages"] if row["id"] == "s000002")
    corrected = _edit_workbook(
        result["workbook"],
        tmp_path,
        sheet="Ranked passages",
        record_id=passage_id,
        edits={"outcome": "miss", "action": "replace", "replacement": "Corrected text"},
    )

    workflow.import_review(run, corrected, detector)

    assert len(detector.calls) >= 2  # prepare, then a full correction rescan
    assert [segment["id"] for segment in detector.calls[1]] == ["s000001", "s000002"]


def test_old_signed_workbook_is_rejected_after_revision_changes(tmp_path: Path):
    def finder(segments):
        return [_finding(segments[0], 0, 13)]

    run, _, result = _prepare(tmp_path, ["Anna Beispiel"], finder)
    old = _edit_workbook(result["workbook"], tmp_path, signoff="First reviewer")
    workflow.import_review(run, old)
    pointer = (run / "current.json").read_bytes()

    stale = _edit_workbook(result["workbook"], tmp_path, signoff="Stale reviewer")
    with pytest.raises(AnonymizerError):
        workflow.import_review(run, stale)
    assert (run / "current.json").read_bytes() == pointer
