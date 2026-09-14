"""Exercise the actual Tk entrypoint and widgets, not just the controller."""

from __future__ import annotations

import os
import select
import shutil
import subprocess
import sys
import time
from pathlib import Path

import pytest


@pytest.fixture(scope="module")
def tk_display():
    pytest.importorskip("tkinter")
    if sys.platform != "linux" or os.environ.get("DISPLAY"):
        yield
        return
    executable = shutil.which("Xvfb")
    if executable is None:
        pytest.skip("Actual Tk tests require a display or Xvfb")
    reader, writer = os.pipe()
    server = subprocess.Popen(
        [executable, "-displayfd", str(writer), "-screen", "0", "1280x800x24", "-nolisten", "tcp"],
        pass_fds=(writer,), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
    os.close(writer)
    try:
        assert select.select([reader], [], [], 10)[0], "Xvfb did not provide a display"
        display = os.read(reader, 32).decode().strip()
        assert display.isdigit(), "Xvfb failed to start"
        with pytest.MonkeyPatch.context() as environment:
            environment.setenv("DISPLAY", f":{display}")
            yield
    finally:
        os.close(reader)
        server.terminate()
        try:
            server.wait(timeout=5)
        except subprocess.TimeoutExpired:
            server.kill()
            server.wait(timeout=5)


def test_real_desktop_entrypoint_enters_event_loop(tk_display, monkeypatch):
    import tkinter as tk

    from transcript_anonymizer import desktop

    factory = tk.Tk
    roots = []
    callbacks = []

    def auto_closing_root():
        root = factory()
        roots.append(root)

        def close():
            callbacks.append("event-loop-running")
            root.destroy()

        root.after(150, close)
        return root

    monkeypatch.setattr(tk, "Tk", auto_closing_root)
    try:
        assert desktop.main([]) == 0
        assert callbacks == ["event-loop-running"]
    finally:
        for root in roots:
            try:
                root.destroy()
            except tk.TclError:
                pass


def _pump(root, predicate, timeout=5):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        root.update()
        if predicate():
            return
        time.sleep(0.01)
    raise AssertionError("Desktop did not reach the expected state")


def test_actual_widgets_display_eta_range_and_elapsed_units(tk_display, monkeypatch):
    import tkinter as tk

    from transcript_anonymizer import desktop

    root = tk.Tk()
    controller = desktop.DesktopController()
    app = desktop.DesktopApp(root=root, controller=controller)
    try:
        app._ui_started_at = time.monotonic() - 3661
        controller.events.put({'kind': 'progress', 'stage': 'processing',
                               'completed': 10, 'total': 100, 'eta': (180, 300)})
        app._drain()
        assert app.eta_text.get() == 'Processing: About 3-5 min remaining'
        controller.events.put({'kind': 'finished', 'success': False})
        app._drain()
        assert app.elapsed_text.get() == 'Elapsed 1:01:01'
        assert app.eta_text.get() == ''
    finally:
        root.destroy()


def test_real_widgets_prepare_review_import_export(tk_display, monkeypatch, tmp_path):
    import tkinter as tk
    from tkinter import filedialog

    from openpyxl import load_workbook

    from transcript_anonymizer import desktop
    from transcript_anonymizer.detection import Detector
    from transcript_anonymizer.documents import read_docx, write_docx

    source = tmp_path / "Synthetic input ä.docx"
    write_docx([{"id": "s000001", "text": "Anna Beispiel: anna@example.invalid"}], source)

    def rules_factory(policy, model, progress):
        detector = Detector(policy, rules_only=True)
        detector.progress = lambda done, total: progress("processing", done, total)
        return detector

    root = tk.Tk()
    errors = []
    root.report_callback_exception = lambda *args: errors.append(args)
    controller = desktop.DesktopController(detector_factory=rules_factory)
    app = desktop.DesktopApp(root=root, controller=controller)
    try:
        monkeypatch.setattr(filedialog, "askopenfilename", lambda **kwargs: str(source))
        app._browse_source()
        assert app.output.get() == str(tmp_path / "Transcript review")
        assert app.export_button.instate(["disabled"])
        app.process_button.invoke()
        assert all(control.instate(["disabled"]) for control in app._controls)
        _pump(root, lambda: app.status_text.get() == "Status: Pending review"
              and app.process_button.instate(["!disabled"]))
        assert "Signed off by" in app.next_step_text.get()
        assert app.open_workbook_button.instate(["!disabled"])
        assert app.export_button.instate(["disabled"])
        assert float(app.progress["value"]) == 100

        workbook = Path(controller.status["workbook"])
        book = load_workbook(workbook)
        sheet = book["Replacements"]
        signoff = next(cell for row in sheet for cell in row if cell.value == "Signed off by")
        sheet.cell(signoff.row, signoff.column + 1).value = "Synthetic reviewer"
        signed = tmp_path / "Reviewed workbook ü.xlsx"
        book.save(signed)
        book.close()
        monkeypatch.setattr(filedialog, "askopenfilename", lambda **kwargs: str(signed))
        app.import_button.invoke()
        _pump(root, lambda: app.export_button.instate(["!disabled"]))
        assert app.status_text.get() == "Status: Signed off"
        app.export_button.invoke()
        _pump(root, lambda: app.open_export_button.instate(["!disabled"]))
        output = controller.export_path
        assert {path.name for path in output.iterdir()} == {"transcript.docx", "manifest.json"}
        text = " ".join(row["text"] for row in read_docx(output / "transcript.docx")["segments"])
        assert "Anna Beispiel" not in text and "anna@example.invalid" not in text
        opened = []
        monkeypatch.setattr(desktop, "_open_local", opened.append)
        app.open_export_button.invoke()
        assert opened == [output]
        assert not errors
    finally:
        controller.cancel()
        controller.join(2)
        root.destroy()


@pytest.mark.parametrize('browse', [True, False])
def test_new_source_clears_previous_review_and_queued_status(tk_display, monkeypatch, tmp_path, browse):
    import tkinter as tk
    from tkinter import filedialog

    from transcript_anonymizer import desktop

    root = tk.Tk()
    controller = desktop.DesktopController()
    app = desktop.DesktopApp(root=root, controller=controller)
    try:
        app.source.set(str(tmp_path / 'first.docx'))
        old = {'run': str(tmp_path / 'old-run'), 'signed_off': True}
        controller._run = tmp_path / 'old-run'
        controller._status = old
        controller._export_path = tmp_path / 'old-export'
        app._apply_status(old)
        app.elapsed_text.set('Elapsed 0:04:20')
        controller.events.put({'kind': 'result', 'action': 'process', 'result': old})
        new = str(tmp_path / 'second.docx')
        if browse:
            monkeypatch.setattr(filedialog, 'askopenfilename', lambda **kwargs: new)
            app._browse_source()
        else:
            app.source.set(new)
        app._drain()
        assert controller.status is None and controller.run_path is None
        assert controller.export_path is None
        assert app.status_text.get() == 'No run selected'
        assert app.elapsed_text.get() == '' and app.eta_text.get() == ''
        assert float(app.progress['value']) == 0
        assert all(button.instate(['disabled']) for button in app._action_buttons)
        assert 'Prepare review' in app.next_step_text.get()
    finally:
        root.destroy()


def test_processing_replaces_loading_instructions(tk_display):
    import tkinter as tk

    from transcript_anonymizer import desktop
    root = tk.Tk()
    controller = desktop.DesktopController()
    app = desktop.DesktopApp(root=root, controller=controller)
    try:
        controller.events.put({'kind': 'progress', 'stage': 'model_loading'})
        app._drain()
        assert 'model loads' in app.next_step_text.get()
        controller.events.put({'kind': 'progress', 'stage': 'processing', 'completed': 0, 'total': 10, 'eta': None})
        app._drain()
        assert 'model loads' not in app.next_step_text.get()
        assert 'processing' in app.next_step_text.get()
        assert app.status_text.get() == 'Status: Processing'
    finally:
        root.destroy()
