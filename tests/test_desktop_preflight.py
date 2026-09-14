"""Headless desktop preflight checks for Windows-like local failures."""

from __future__ import annotations

import time
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from transcript_anonymizer import workflow
from transcript_anonymizer.desktop import DesktopController
from transcript_anonymizer.detection import Detector
from transcript_anonymizer.documents import read_docx, write_docx


def _finish(controller: DesktopController, timeout: float = 5.0) -> list[dict]:
    deadline = time.monotonic() + timeout
    messages: list[dict] = []
    while time.monotonic() < deadline:
        messages.extend(controller.messages())
        if any(message["kind"] == "finished" for message in messages):
            controller.join(0.1)
            messages.extend(controller.messages())
            return messages
        time.sleep(0.01)
    raise AssertionError(f"controller did not finish: {messages}")


def _source(tmp_path: Path, text: str = "Anna Beispiel: anna@example.invalid") -> Path:
    source = tmp_path / "Eingabe mit Leerzeichen-ä.docx"
    write_docx([{"id": "s000001", "text": text}], source)
    return source


def _rules_factory(policy, _model_path, progress):
    detector = Detector(policy, rules_only=True)
    detector.progress = lambda completed, total: progress("processing", completed, total)
    return detector


def _controller(**kwargs) -> DesktopController:
    return DesktopController(detector_factory=_rules_factory, **kwargs)


def _sign(workbook: Path, tmp_path: Path, name: str = "Synthetic reviewer") -> Path:
    book = load_workbook(workbook)
    ws = book["Replacements"]
    label = next(cell for row in ws for cell in row if cell.value == "Signed off by")
    ws.cell(label.row, label.column + 1).value = name
    output = tmp_path / f"signed-{len(list(tmp_path.glob('signed-*.xlsx')))}.xlsx"
    book.save(output)
    book.close()
    return output


def _edit_group(workbook: Path, tmp_path: Path, group_id: str) -> Path:
    book = load_workbook(workbook)
    ws = book["Replacements"]
    header = next(row for row in ws if row[0].value == "id")
    columns = {cell.value: cell.column for cell in header}
    row = next(row for row in ws if row[0].value == group_id)
    ws.cell(row[0].row, columns["action"]).value = "replace"
    ws.cell(row[0].row, columns["replacement"]).value = "PERSON_CUSTOM"
    output = tmp_path / f"changed-{len(list(tmp_path.glob('changed-*.xlsx')))}.xlsx"
    book.save(output)
    book.close()
    return output


def _prepare(controller: DesktopController, tmp_path: Path):
    source = _source(tmp_path)
    controller.start_process(source, tmp_path / "Review folder ä")
    messages = _finish(controller)
    return source, next(message["result"] for message in messages if message["kind"] == "result")


def test_unicode_space_paths_work_end_to_end_without_tk(tmp_path: Path):
    controller = _controller()
    source, result = _prepare(controller, tmp_path)
    assert source.name.endswith("ä.docx")
    assert result["workbook"].endswith("review.xlsx")

    controller.import_workbook(_sign(Path(result["workbook"]), tmp_path))
    imported = next(message["result"] for message in _finish(controller) if message["kind"] == "result")
    assert imported["signed_off"] is True
    controller.export(tmp_path / "Export Ziel ü")
    assert any(message["kind"] == "result" for message in _finish(controller))
    exported = controller.export_path / "transcript.docx"
    assert exported.is_file()
    assert "Anna Beispiel" not in " ".join(
        segment["text"] for segment in read_docx(exported)["segments"]
    )


def test_corrupt_docx_is_error_and_does_not_create_success_state(tmp_path: Path):
    controller = _controller()
    source = tmp_path / "corrupt input.docx"
    source.write_bytes(b"not a docx")
    controller.start_process(source, tmp_path / "reviews")
    messages = _finish(controller)
    assert any(message["kind"] == "error" for message in messages)
    assert not any(message["kind"] == "result" for message in messages)
    assert controller.status is None
    assert controller.run_path is None


def test_corrupt_workbook_is_rejected_without_revision_change(tmp_path: Path):
    controller = _controller()
    _, result = _prepare(controller, tmp_path)
    run = controller.run_path
    before = (run / "current.json").read_bytes()
    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_bytes(b"not an xlsx")

    controller.import_workbook(corrupt)
    messages = _finish(controller)
    assert any(message["kind"] == "error" for message in messages)
    assert (run / "current.json").read_bytes() == before
    assert result["signed_off"] is False


def test_disk_full_and_unwritable_output_are_recoverable_errors(tmp_path: Path):
    controller = _controller()
    source = _source(tmp_path)
    with patch.object(workflow, "write_workbook", side_effect=OSError("disk full")):
        controller.start_process(source, tmp_path / "disk full review")
        messages = _finish(controller)
    assert any(message["kind"] == "error" for message in messages)
    assert controller.status is None
    assert controller.run_path is None

    occupied = tmp_path / "not-a-folder"
    occupied.write_text("occupied", encoding="utf-8")
    controller.start_process(source, occupied)
    messages = _finish(controller)
    assert any(message["kind"] == "error" for message in messages)
    assert controller.status is None


def test_missing_model_fails_closed_in_default_controller(tmp_path: Path):
    controller = DesktopController(model_path=tmp_path / "missing-model")
    controller.start_process(_source(tmp_path), tmp_path / "reviews")
    messages = _finish(controller)
    assert any(message["kind"] == "error" for message in messages)
    assert not any(message["kind"] == "result" for message in messages)
    assert controller.status is None


def test_stale_workbook_and_tampered_candidate_are_blocked(tmp_path: Path):
    controller = _controller()
    _, result = _prepare(controller, tmp_path)
    old = _sign(Path(result["workbook"]), tmp_path, "First reviewer")
    controller.import_workbook(old)
    next(message["result"] for message in _finish(controller) if message["kind"] == "result")
    run = controller.run_path
    pointer = (run / "current.json").read_bytes()

    stale = _sign(Path(result["workbook"]), tmp_path, "Stale reviewer")
    controller.import_workbook(stale)
    assert any(message["kind"] == "error" for message in _finish(controller))
    assert (run / "current.json").read_bytes() == pointer

    _state, folder = workflow._load(run)
    with (folder / "candidate.docx").open("ab") as handle:
        handle.write(b"tamper")
    controller.select_run(run)
    messages = _finish(controller)
    assert any(message["kind"] == "error" for message in messages)
    assert controller.status is None


def test_changed_import_clears_previous_export_path(tmp_path: Path):
    controller = _controller()
    _, result = _prepare(controller, tmp_path)
    controller.import_workbook(_sign(Path(result["workbook"]), tmp_path))
    signed = next(message["result"] for message in _finish(controller) if message["kind"] == "result")
    controller.export(tmp_path / "exports")
    _finish(controller)
    assert controller.export_path is not None

    state, _ = workflow._load(controller.run_path)
    group_id = state["view"]["groups"][0]["id"]
    controller.import_workbook(_edit_group(Path(signed["workbook"]), tmp_path, group_id))
    changed = next(message["result"] for message in _finish(controller) if message["kind"] == "result")
    assert changed["signed_off"] is False
    assert controller.export_path is None
