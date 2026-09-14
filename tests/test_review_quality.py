from pathlib import Path
from typing import ClassVar

import pytest
from openpyxl import load_workbook

from transcript_anonymizer import workflow
from transcript_anonymizer.documents import write_docx
from transcript_anonymizer.errors import AnonymizerError
from transcript_anonymizer.ranking import rank_retained_passages, resolve_findings


def finding(segment_id, text, start, end, category="person", score=0.6, review_only=False):
    return {
        "id": f"{segment_id}-{start}-{end}-{category}",
        "segment_id": segment_id,
        "start": start,
        "end": end,
        "text": text[start:end],
        "category": category,
        "score": score,
        "detector": "fake",
        **({"review_only": True} if review_only else {}),
    }


def test_overlap_resolution_preserves_connected_source_coverage():
    text = "Anna Beispiel sprach."
    selected, suppressed = resolve_findings(
        [
            finding("s1", text, 0, 4, score=0.95),
            finding("s1", text, 0, 13, score=0.60),
        ]
    )
    assert [(row["start"], row["end"], row["text"]) for row in selected] == [
        (0, 13, "Anna Beispiel")
    ]
    assert suppressed


def test_overlap_resolution_merges_transitive_overlap_chain_and_provenance():
    text = "Anna Beispiel"
    selected, suppressed = resolve_findings(
        [
            finding("s1", text, 0, 5, score=0.7) | {"detector": "first"},
            finding("s1", text, 4, 8, score=0.8) | {"detector": "second"},
            finding("s1", text, 7, len(text), score=0.6) | {"detector": "third"},
        ]
    )
    assert [(row["start"], row["end"], row["text"]) for row in selected] == [
        (0, len(text), text)
    ]
    assert selected[0]["provenance"]["detectors"] == ["first", "second", "third"]
    assert {row["reason"] for row in suppressed} == {"overlap_merged"}


def test_overlap_resolution_leaves_disjoint_and_review_only_findings_unchanged():
    text = "Anna Example spoke."
    disjoint = finding("s1", text, 0, 4, score=0.95)
    review = finding("s1", text, 5, 12, score=0.4, review_only=True)
    selected, suppressed = resolve_findings([disjoint, review])
    assert selected == [disjoint]
    assert suppressed == []


def test_overlap_resolution_omits_ambiguous_entity_key():
    text = "Anna Beispiel"
    first = finding("s1", text, 0, 4) | {"entity_key": "anna"}
    second = finding("s1", text, 0, len(text), score=0.5) | {"entity_key": "beispiel"}
    selected, _ = resolve_findings([first, second])
    assert "entity_key" not in selected[0]


def test_overlap_resolution_preserves_full_span_explicit_identity():
    text = "Anna Beispiel"
    full = finding("s1", text, 0, len(text)) | {"entity_key": "anna"}
    short = finding("s1", text, 0, 4, score=0.99)
    selected, _ = resolve_findings([full, short])
    assert selected[0]["entity_key"] == "anna"


def test_overlap_resolution_does_not_extend_partial_identity_claim():
    text = "Anna Beispiel"
    full = finding("s1", text, 0, len(text))
    short = finding("s1", text, 0, 4, score=0.99) | {"entity_key": "anna"}
    selected, _ = resolve_findings([full, short])
    assert "entity_key" not in selected[0]


def test_reconcile_without_resolver_version_keeps_legacy_projection():
    text = "Anna Beispiel"
    state = {
        "segments": [{"id": "s1", "text": text}],
        "findings": [
            finding("s1", text, 0, 4, score=0.95),
            finding("s1", text, 0, len(text), score=0.6),
        ],
        "groups": {},
        "token_counters": {},
        "occurrence_decisions": {},
    }
    workflow._reconcile(state)
    assert [(row["start"], row["end"]) for row in state["selected"]] == [(0, 4)]


def test_review_only_finding_is_ranked_but_not_replacement_group():
    segments = [{"id": "s1", "text": "Alice spoke."}]
    state = {
        "segments": segments,
        "findings": [finding("s1", segments[0]["text"], 0, 5, review_only=True)],
        "groups": {},
        "token_counters": {},
        "occurrence_decisions": {},
    }
    workflow._reconcile(state)
    assert state["selected"] == []
    assert state["groups"] == {}
    assert state["suppressed"][0]["reason"] == "review_only"

    rendered = [{"id": "s1", "text": "Alice spoke."}]
    rows = rank_retained_passages(segments, rendered, state["findings"], [])
    assert "review_only_candidate" in rows[0]["reasons"]


def test_masked_findings_do_not_boost_visible_candidate():
    source = [{"id": "s1", "text": "Alice spoke."}]
    rendered = [{"id": "s1", "text": "<REDACTED> spoke."}]
    row = rank_retained_passages(
        source,
        rendered,
        [finding("s1", source[0]["text"], 0, 5, score=0.2)],
        [{"segment_id": "s1", "start": 0, "end": 5}],
    )[0]
    assert "unresolved_candidate" not in row["reasons"]
    assert "low_confidence_candidate" not in row["reasons"]
    assert row["score"] == 0


def test_adjacent_masked_intervals_cover_one_finding_without_boost():
    text = "Alice Example"
    row = rank_retained_passages(
        [{"id": "s1", "text": text}],
        [{"id": "s1", "text": "<REDACTED>"}],
        [finding("s1", text, 0, len(text), score=0.2)],
        [
            {"segment_id": "s1", "start": 0, "end": 5},
            {"segment_id": "s1", "start": 5, "end": len(text)},
        ],
    )[0]
    assert "unresolved_candidate" not in row["reasons"]


def test_whitespace_gap_between_masks_is_not_visible_evidence():
    text = "Alice  Example"
    row = rank_retained_passages(
        [{"id": "s1", "text": text}],
        [{"id": "s1", "text": "<REDACTED>"}],
        [finding("s1", text, 0, len(text), category="full_name", score=0.2)],
        [
            {"segment_id": "s1", "start": 0, "end": 5},
            {"segment_id": "s1", "start": 7, "end": len(text)},
        ],
    )[0]
    assert "unresolved_candidate" not in row["reasons"]


def test_duplicate_evidence_does_not_double_boost_a_passage():
    text = "Alice spoke."
    one = rank_retained_passages(
        [{"id": "s1", "text": text}],
        [{"id": "s1", "text": text}],
        [finding("s1", text, 0, 5, score=0.6)],
        [],
    )[0]
    duplicate = rank_retained_passages(
        [{"id": "s1", "text": text}],
        [{"id": "s1", "text": text}],
        [finding("s1", text, 0, 5, score=0.6), finding("s1", text, 0, 5, score=0.6)],
        [],
    )[0]
    assert duplicate["score"] == one["score"]


def test_partial_mask_and_overlap_evidence_are_ranked():
    text = "Alice Example spoke."
    source = [{"id": "s1", "text": text}]
    rendered = [{"id": "s1", "text": "<REDACTED> Example spoke."}]
    rows = rank_retained_passages(
        source,
        rendered,
        [
            finding("s1", text, 0, 5, score=0.9),
            finding("s1", text, 0, 13, category="full_name", score=0.4),
        ],
        [{"segment_id": "s1", "start": 0, "end": 5}],
    )
    reasons = rows[0]["reasons"]
    assert "overlap_suppressed_candidate" in reasons
    assert "low_confidence_candidate" in reasons


def test_remainder_sample_is_deterministic_and_bounded():
    segments = [{"id": f"s{i}", "text": "ordinary process note"} for i in range(10)]
    first = rank_retained_passages(segments, segments, [], [], sample_size=3, sample_seed="fixed")
    second = rank_retained_passages(segments, segments, [], [], sample_size=3, sample_seed="fixed")
    assert [row["id"] for row in first] == [row["id"] for row in second]
    assert sum("remainder_sample" in row["reasons"] for row in first) == 3


def test_prepare_can_opt_out_of_remainder_sample(tmp_path: Path):
    detector = ContextDetector()
    source = tmp_path / "source.docx"
    write_docx(
        [{"id": "s000001", "text": "ordinary note", "locator": {"kind": "paragraph", "index": 0}}],
        source,
    )
    result = workflow.prepare(source, tmp_path / "run", {}, detector, review_sample_size=0)
    state, _ = workflow._load(tmp_path / "run")
    assert all("remainder_sample" not in row["reasons"] for row in state["view"]["passages"])
    assert result["sample_recommended"] == 0


class ContextDetector:
    metadata: ClassVar[dict] = {"detector": "context-fake", "version": 1}

    def __init__(self):
        self.calls = []

    def detect(self, segments):
        self.calls.append([segment["id"] for segment in segments])
        return [
            finding(segment["id"], segment["text"], 0, 5, score=0.9)
            for segment in segments
            if segment["text"].startswith("Alice")
        ]


class OverlapDetector:
    metadata: ClassVar[dict] = {"detector": "overlap-fake", "version": 1}

    def detect(self, segments):
        return [
            finding(segment["id"], segment["text"], 0, 4, score=0.95) | {"detector": "short"}
            for segment in segments
            if segment["text"].startswith("Anna")
        ] + [
            finding(segment["id"], segment["text"], 0, 13, score=0.60) | {"detector": "full"}
            for segment in segments
            if segment["text"].startswith("Anna")
        ]


def test_prepare_uses_coverage_preserving_overlap_resolver(tmp_path: Path):
    source = tmp_path / "source.docx"
    write_docx(
        [
            {
                "id": "s000001",
                "text": "Anna Beispiel sprach.",
                "locator": {"kind": "paragraph", "index": 0},
            }
        ],
        source,
    )
    run = tmp_path / "run"
    workflow.prepare(source, run, {}, OverlapDetector())
    state, _ = workflow._load(run)
    assert state["overlap_resolver_version"] == 2
    assert [(row["start"], row["end"], row["text"]) for row in state["selected"]] == [
        (0, 13, "Anna Beispiel")
    ]
    assert workflow._render(state)[0]["text"] == "PERSON_001 sprach."


def test_legacy_saved_run_cannot_be_imported_or_exported(tmp_path: Path):
    source = tmp_path / "source.docx"
    write_docx(
        [{"id": "s000001", "text": "Anna spoke.", "locator": {"kind": "paragraph", "index": 0}}],
        source,
    )
    run = tmp_path / "run"
    result = workflow.prepare(source, run, {}, ContextDetector())
    state, folder = workflow._load(run)
    state.pop("overlap_resolver_version")
    raw = workflow.encoded(state)
    (folder / "state.json").write_bytes(raw)
    (run / "current.json").write_bytes(
        workflow.encoded({"directory": folder.name, "sha256": workflow.digest(raw)})
    )
    assert workflow.status(run)["overlap_resolver_version"] == 1
    with pytest.raises(AnonymizerError, match="prepare a new run"):
        workflow.import_review(run, Path(result["workbook"]), ContextDetector())
    with pytest.raises(AnonymizerError, match="prepare a new run"):
        workflow.export(run, tmp_path / "export")


def test_passage_correction_rescans_full_working_document(tmp_path: Path):
    detector = ContextDetector()
    source = tmp_path / "source.docx"
    write_docx(
        [
            {"id": "s000001", "text": "Alice spoke.", "locator": {"kind": "paragraph", "index": 0}},
            {
                "id": "s000002",
                "text": "Neighbor context.",
                "locator": {"kind": "paragraph", "index": 1},
            },
        ],
        source,
    )
    run = tmp_path / "run"
    result = workflow.prepare(source, run, {}, detector)
    assert detector.calls[-1] == ["s000001", "s000002"]
    state, _ = workflow._load(run)
    passage_id = state["view"]["passages"][0]["id"]
    book = load_workbook(result["workbook"])
    sheet = book["Ranked passages"]
    header = {
        cell.value: cell.column for cell in next(row for row in sheet if row[0].value == "id")
    }
    row = next(row for row in sheet if row[0].value == passage_id)[0].row
    sheet.cell(row, header["action"]).value = "replace"
    sheet.cell(row, header["replacement"]).value = "clean text"
    edited = tmp_path / "correction.xlsx"
    book.save(edited)
    book.close()
    workflow.import_review(run, edited, detector)
    assert detector.calls[-1] == ["s000001", "s000002"]
